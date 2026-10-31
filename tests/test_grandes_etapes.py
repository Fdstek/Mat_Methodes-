"""Grandes étapes (ex. Degumming) contenant des sous-étapes, numérotation continue."""
import json

import pdfplumber

from mat_methode.domain import journal as J
from mat_methode.domain import regles as R
from mat_methode.domain.models import Manip, Reglages
from mat_methode.exports.echange import exporter, importer
from mat_methode.exports.pdf import ecrire_pdf
from mat_methode.exports.word import ecrire_docx
from mat_methode.report.html import rapport_html
from mat_methode.report.rapport import construire, markdown


def degumming(m: Manip):
    """Chauffage, puis Degumming (acide, agitation, centrifugation), puis séchage."""
    R.ajouter_etape(m, "chauffage")
    g = R.creer_groupe(m, "Degumming")
    for op in ("ajout d’acide phosphorique", "agitation", "centrifugation"):
        R.ajouter_etape(m, op, groupe_id=g.id)
    R.ajouter_etape(m, "séchage")
    return g


def test_blocs_et_plage():
    m = Manip()
    g = degumming(m)
    assert [(x.nom if x else None, idx) for x, idx in R.blocs(m)] == [(None, [0]), ("Degumming", [1, 2, 3]), (None, [4])]
    assert R.numeros_groupe(m, g.id) == [2, 3, 4]
    assert (R.plage([2, 3, 4]), R.plage([3]), R.plage([2, 5])) == ("étapes 2 à 4", "étape 3", "étapes 2, 5")


def test_numerotation_continue_et_etiquettes():
    m = Manip()
    g = degumming(m)
    centri = m.etapes[3]
    p = R.nouvel_echantillon(m, Reglages(), centri.id)
    assert R.etiquettes(m)[p.id] == "S4"  # étiquette inchangée : numéro continu
    R.ajouter_etape(m, "lavage", apres=1)  # inséré au milieu de la grande étape : en fait partie
    assert m.etapes[2].groupe_id == g.id and R.numeros_groupe(m, g.id) == [2, 3, 4, 5]
    R.ajouter_etape(m, "broyage", apres=-1)
    assert R.numeros_groupe(m, g.id) == [3, 4, 5, 6] and R.etiquettes(m)[p.id] == "S6"


def test_placer_defaire_deplacer():
    m = Manip()
    g = degumming(m)
    sechage = m.etapes[4]
    R.placer_dans_groupe(m, sechage.id, g.id)  # déjà voisine : rejoint la grande étape sans bouger
    assert R.numeros_groupe(m, g.id) == [2, 3, 4, 5]
    chauffage = m.etapes[0]
    R.placer_dans_groupe(m, chauffage.id, None)
    R.ajouter_etape(m, "filtration")
    R.placer_dans_groupe(m, m.etapes[-1].id, g.id)
    assert R.numeros_groupe(m, g.id) == [2, 3, 4, 5, 6]
    assert R.deplacer_groupe(m, g.id, -1) and m.etapes[0].groupe_id == g.id and m.etapes[-1] is chauffage
    assert not R.deplacer_groupe(m, g.id, -1)
    R.defaire_groupe(m, g.id)
    assert m.groupes == [] and all(e.groupe_id is None for e in m.etapes) and len(m.etapes) == 6


def test_texte_du_rapport():
    m = Manip()
    degumming(m)
    r = construire(m, Reglages())
    assert "Grandes étapes : Degumming (étapes 2 à 4)." in r.sections[0].texte
    t = r.sections[2].texte
    assert "Étape 1 : le produit a été chauffé" in t and "Degumming (étapes 2 à 4). Étape 2 :" in t and "Étape 5 :" in t
    ph = r.phases
    assert [p.groupe for p in ph] == ["", "", "Degumming", "Degumming", "Degumming", ""]
    assert [p.debut_groupe for p in ph] == [False, False, True, False, False, False]
    md = markdown(m, Reglages())
    assert "**Grande étape : Degumming (étapes 2 à 4)**\n\n#### 2 · Ajout d’acide phosphorique" in md
    assert "Grande étape : Degumming (étapes 2 à 4)" in rapport_html(m, Reglages()) and "DEGUMMING" in rapport_html(m, Reglages())


def test_duplication_copie_les_grandes_etapes():
    m = Manip()
    g = degumming(m)
    c = R.dupliquer(m, Reglages())
    assert len(c.groupes) == 1 and c.groupes[0].id != g.id and c.groupes[0].nom == "Degumming"
    assert [e.groupe_id for e in c.etapes] == [None] + [c.groupes[0].id] * 3 + [None]


def test_journal_apres_lancement():
    m = Manip()
    g = degumming(m)
    J.lancer(m, par="Pierre")
    g.nom = "Dégommage acide"
    R.placer_dans_groupe(m, m.etapes[4].id, g.id)
    R.ajouter_etape(m, "lavage", groupe_id=g.id, apres=4)
    txt = [x.txt for x in J.journaliser(m)]
    assert "Grande étape renommée : Degumming → Dégommage acide" in txt
    assert "Étape 5 · Séchage : grande étape aucune → Dégommage acide" in txt
    assert "Étape ajoutée : Étape 6 · Lavage (dans Dégommage acide)" in txt


def test_echange_aller_retour():
    m = Manip(nom="Essai")
    degumming(m)
    r = Reglages(grandes_etapes=[])
    d = json.loads(json.dumps(exporter(m, r)))
    assert d["grandes_etapes"] == [{"nom": "Degumming", "etapes": [2, 3, 4]}]
    assert [e.get("groupe") for e in d["etapes"]] == [None, "Degumming", "Degumming", "Degumming", None]
    m2, _ = importer(d, r)
    assert [(x.nom if x else None, idx) for x, idx in R.blocs(m2)] == [(None, [0]), ("Degumming", [1, 2, 3]), (None, [4])]
    assert "Degumming" in r.grandes_etapes


def test_sans_grande_etape_format_inchange(ethanol, reglages):
    d = exporter(ethanol, reglages)
    assert "grandes_etapes" not in d and all("groupe" not in e for e in d["etapes"])


def test_pdf_et_word(tmp_path):
    import docx
    m = Manip(nom="Essai")
    degumming(m)
    d = exporter(m, Reglages())
    ecrire_pdf(d, tmp_path / "a.pdf")
    with pdfplumber.open(tmp_path / "a.pdf") as pdf:
        texte = " ".join(" ".join(p.extract_text() for p in pdf.pages).split())
    assert "DEGUMMING" in texte and "Grandes étapes (bandeau en haut des cartes) : Degumming (étapes 2 à 4)" in texte
    ecrire_docx(d, tmp_path / "a.docx")
    w = docx.Document(str(tmp_path / "a.docx"))
    tout = "\n".join(p.text for p in w.paragraphs) + "\n".join(c.text for t in w.tables for r in t.rows for c in r.cells)
    assert "Grandes étapes : Degumming (étapes 2 à 4)" in tout and "DEGUMMING · 2 · Ajout d’acide phosphorique" in tout


def test_reglages_par_defaut_et_fusion():
    from mat_methode.storage.fusion import fusionner
    base = Reglages()
    assert base.grandes_etapes[0] == "Degumming"
    mien, leur = base.model_copy(deep=True), base.model_copy(deep=True)
    mien.grandes_etapes.append("Estérification")
    leur.grandes_etapes.append("Hydrogénation")
    assert {"Estérification", "Hydrogénation"} <= set(fusionner(base, mien, leur).grandes_etapes)
