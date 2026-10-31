"""Lecture du fichier Scirius (même logique que le projet Fractionnement), sur un faux classeur."""
from datetime import datetime

import pytest
from openpyxl import Workbook

from mat_methode import scirius as S

ENTETE = ["Date réception", "Fournisseur", "Product Name souhaité", "Stock Item", "Stockage", "Nb", "Quantité",
          "Batch/Lot", "Nom du produit", "Destinataire", "Équipe"]


def classeur(chemin, entete=ENTETE, ligne=8):
    wb = Workbook()
    ws = wb.active
    ws.title = "Suivi General"
    ws.cell(1, 1, "KS Scirius - suivi")
    for i, h in enumerate(entete, start=1):
        ws.cell(ligne, i, h)
    lignes = [
        [datetime(2026, 9, 12), "Sigma", "", "001245", "Armoire inflammables 1", 1, "2,5 L", "SHBQ1234", "Éthanol 96 %", "P.BURGUN", "Bio"],
        [datetime(2026, 10, 2), "Interne", "", "000777", "Chambre froide", 1, "20 kg", "GA-2026-09", "Graisse animale brute", "P.BURGUN", ""],
        [datetime(2026, 10, 20), "Merck", "", "001302", "Armoire inflammables 2", 2, "1 L", "K55", "Acétone", "C.LACOUX", ""],
    ]
    for k, ln in enumerate(lignes, start=ligne + 1):
        for i, v in enumerate(ln, start=1):
            ws.cell(k, i, v)
    wb.save(chemin)
    return chemin


def test_inspecter_et_lire(tmp_path):
    info = S.inspecter(classeur(tmp_path / S.FICHIER))
    assert info["ok"] and info["entete"] == 8
    lignes = S.lire(info)
    assert len(lignes) == 3 and lignes[0]["stock_item"] == "001245" and lignes[0]["date"] == "12/09/2026"
    assert lignes[1]["product"] == "Graisse animale brute" and lignes[1]["batch"] == "GA-2026-09"


def test_entete_deplacee_et_mauvais_fichier(tmp_path):
    assert S.inspecter(classeur(tmp_path / "a.xlsx", ligne=5))["entete"] == 5
    mauvais = classeur(tmp_path / "b.xlsx", entete=["Date", "Nom"])
    info = S.inspecter(mauvais)
    assert not info["ok"] and "colonnes manquantes" in info["raison"]


def test_chercher(tmp_path):
    lignes = S.lire(S.inspecter(classeur(tmp_path / S.FICHIER)))
    assert [x["stock_item"] for x in S.chercher(lignes, "")] == ["001302", "000777", "001245"]  # plus récent d'abord
    assert [x["product"] for x in S.chercher(lignes, "ethanol sigma")] == ["Éthanol 96 %"]
    assert [x["product"] for x in S.chercher(lignes, "GA-2026")] == ["Graisse animale brute"]


def test_source_synchronisee_puis_cache(tmp_path):
    racine = tmp_path / "OneDrive - TotalEnergies" / "Magasin" / "Suivi enregistrement échantillons"
    racine.mkdir(parents=True)
    classeur(racine / S.FICHIER)
    appels = []
    src = S.SourceScirius(tmp_path / "cache", utilisateur="J0000001", racines=[tmp_path / "OneDrive - TotalEnergies"],
                          excel=lambda url, cible: (appels.append(url), (None, "pas d’Excel"))[1])
    lu = src.lire()
    assert len(lu.lignes) == 3 and lu.origine.startswith("fichier synchronisé") and not appels
    assert src.copie.exists() and src.f_lignes.exists()
    # deuxième lecture : depuis le cache (le fichier n'a pas changé)
    assert "lu le" in src.lire().origine and len(src.lire().lignes) == 3


def test_source_introuvable_puis_copie_partagee(tmp_path):
    src = S.SourceScirius(tmp_path / "cache", racines=[], excel=lambda url, cible: (None, "pas d’Excel"))
    lu = src.lire()
    assert lu.lignes == [] and any("Excel" in e for e in lu.ecartes)
    (tmp_path / "cache").mkdir(exist_ok=True)
    classeur(src.copie)
    lu = src.lire(forcer=True)
    assert len(lu.lignes) == 3 and lu.origine.startswith("copie partagée")


def test_fichier_choisi(tmp_path):
    f = classeur(tmp_path / "mon fichier.xlsx")
    lu = S.SourceScirius(tmp_path / "cache", fichier_choisi=str(f), racines=[], excel=lambda u, c: (None, "")).lire()
    assert lu.origine.startswith("fichier choisi") and len(lu.lignes) == 3


@pytest.mark.parametrize("nom,ini", [("Pierre BURGUN", "PB"), ("Jean-Marc DUPONT", "JMD"), ("Interne · R&D", "IR")])
def test_initiales(nom, ini):
    from mat_methode.domain.models import initiales_de
    assert initiales_de(nom) == ini
