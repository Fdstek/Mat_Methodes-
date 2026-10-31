"""Exports JSON, PDF, Word ; import (critères 9, 10, 15, 17, 18, 20, 23, 32)."""
import json
import zipfile

import pdfplumber
import pytest

from mat_methode.domain.models import Reglages
from mat_methode.exports.echange import ImportInvalide, exporter, importer
from mat_methode.exports.pdf import ecrire_pdf, polices, titre_lignes
from mat_methode.exports.word import ecrire_docx
from mat_methode.extraction.extract_mm import extraire
from mat_methode.report.rapport import markdown

from .conftest import EXEMPLES


def test_json_compatible_avec_la_reference(ethanol, reglages):
    a = exporter(ethanol, reglages)
    b = json.loads((EXEMPLES / "manip-pour-script-pdf.json").read_text(encoding="utf-8"))
    assert set(a) == set(b)
    for k in b:
        if k in ("sous_sections", "etapes"):
            continue
        assert a[k] == b[k], k
    # étapes : mêmes champs, plus « autres » (conditions libres) ; « role » n'existe plus
    for x, y in zip(a["etapes"], b["etapes"]):
        x = {k: v for k, v in x.items() if k != "autres"}
        y = dict(y, reactifs_detail=[{"nom": r["nom"], "quantite": r["quantite"]} for r in y["reactifs_detail"]])
        assert x == y
    assert [s["titre"] for s in a["sous_sections"]] == [s["titre"] for s in b["sous_sections"]]


def test_crit23_aller_retour_json(ethanol, reglages):
    ethanol.contexte.type = "Test d’idée"
    ethanol.contexte.idee = "Idée"
    ethanol.etapes[1].autres = [{"nom": "pH", "valeur": "7"}]
    m, _ = importer(json.loads(json.dumps(exporter(ethanol, reglages))), reglages)
    assert m.nom == ethanol.nom and m.tags == ethanol.tags and m.contexte.idee == "Idée"
    assert m.contexte.objectif == ethanol.contexte.objectif and m.contexte.volume_depart_ml == 1000
    assert len(m.journal) == 2 and m.lancee.par == "Pierre"
    assert m.etapes[1].autres[0].nom == "pH"
    ethanol.rev = m.rev
    a, b = markdown(ethanol, reglages), markdown(m, reglages)
    assert a.split("## 2.")[1] == b.split("## 2.")[1]


def test_crit10_import_extraction():
    r = Reglages(reactifs=[])
    d = json.loads((EXEMPLES / "mm_extrait.json").read_text(encoding="utf-8"))
    m, a_verifier = importer(d, r)
    assert len(m.etapes) == 5 and m.etat == r.etat_initial() and a_verifier == []
    assert [x.nom for x in r.reactifs] == ["Éthanol 96 %", "Eau distillée"]
    assert m.etapes[2].reactifs[0].quantite == "1 : 1 (v/v)"


def test_import_invalide(reglages):
    for d in ({}, {"etapes": "x"}, [], 3):
        with pytest.raises(ImportInvalide):
            importer(d, reglages)


def test_extraction_pdf():
    res = extraire(EXEMPLES / "rapport_exemple.pdf")
    ref = json.loads((EXEMPLES / "mm_extrait.json").read_text(encoding="utf-8"))
    assert res["sous_sections"] == ref["sous_sections"] and res["reactifs"] == ref["reactifs"]
    assert [e["operation"] for e in res["etapes"]] == [e["operation"] for e in ref["etapes"]]
    assert res["etapes"][1]["parametres"] == {"filtre_um": 5.0, "vide": True}  # « sous vide » détecté


def test_crit9_15_pdf(ethanol, reglages, tmp_path):
    f = tmp_path / "manip.pdf"
    ecrire_pdf(exporter(ethanol, reglages), f)
    with pdfplumber.open(f) as pdf:
        texte = " ".join(" ".join(p.extract_text() for p in pdf.pages).split())
        assert 2 <= len(pdf.pages) <= 4
    assert "5 µm" in texte and "60 °C" in texte and "Désacidification" in texte
    for x in ("S0", "S2", "S3a", "S3b", "S5a", "S5b", "Échantillons", "Détail par phase du flux",
              "Modifications après lancement", "Modifiée après lancement (2)", "à saisir"):
        assert x in texte, x


def test_crit20_titres_longs():
    f = polices()[1]
    lignes, taille = titre_lignes("Chauffage préalable de la graisse brute", 54, f)
    assert len(lignes) >= 2 and taille == 8.5 and " ".join(lignes) == "Chauffage préalable de la graisse brute"
    lignes, taille = titre_lignes("Électrocoagulation", 54, f)
    assert taille < 8.5 and taille >= 6.5
    from reportlab.pdfbase.pdfmetrics import stringWidth
    assert all(stringWidth(x, f, taille) <= 54 for x in lignes)


def test_crit17_word(ethanol, reglages, tmp_path):
    import docx
    from docx.enum.text import WD_COLOR_INDEX
    ethanol.etapes[0].operation = "chauffage préalable de la graisse brute"
    f = tmp_path / "manip.docx"
    ecrire_docx(exporter(ethanol, reglages), f)
    d = docx.Document(str(f))
    texte = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    for x in ("Matériel et méthodes", "Détail par phase du flux", "Réactifs", "Échantillons", "S3a", "µm", "°C",
              "Modifications après lancement", "Modifiée après lancement (2)"):
        assert x in texte, x
    surligne = [r for p in d.paragraphs for r in p.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert surligne and all("à saisir" in r.text for r in surligne)
    assert len(d.inline_shapes) == 1
    with zipfile.ZipFile(f) as z:
        assert any(n.startswith("word/media/") for n in z.namelist())
