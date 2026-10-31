"""Export Word (.docx) à partir du JSON d'échange (repris de `scripts/export_word.py`).

Le schéma du flux est une image PNG (Pillow) ; tout le reste est du vrai texte modifiable.
Les « à saisir » sont surlignés en jaune.
"""
from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from mat_methode.domain.texte import fmt_dt, fmt_dur, fmt_vol
from mat_methode.exports.pdf import code_badge, conditions, moment_texte
from mat_methode.exports.polices import trouver

TEAL, AMBER, MUTED = "0F6B5C", "8A4B00", "566360"
REPERE = re.compile(r"(\[[^\]]*à saisir\])")


def police(gras: bool, taille: int):
    paire = trouver()
    if paire:
        return ImageFont.truetype(paire[1] if gras else paire[0], taille)
    return ImageFont.load_default()


def _lignes(draw, texte, font, w):
    out, cur = [], ""
    for mot in texte.split():
        test = (cur + " " + mot).strip()
        if draw.textlength(test, font=font) > w and cur:
            out.append(cur)
            cur = mot
        else:
            cur = test
    return out + ([cur] if cur else [])


def _titre(draw, texte, w):
    """Titre d'étape : police réduite si un mot est trop large (26 -> 18 px), puis coupure du mot."""
    taille = 26
    while taille > 18 and any(draw.textlength(m, font=police(True, taille)) > w for m in texte.split()):
        taille -= 1
    f = police(True, taille)
    mots = []
    for mot in texte.split():
        while draw.textlength(mot, font=f) > w and len(mot) > 2:
            k = len(mot) - 1
            while k > 1 and draw.textlength(mot[:k] + "-", font=f) > w:
                k -= 1
            mots.append(mot[:k] + "-")
            mot = mot[k:]
        mots.append(mot)
    lignes, cur = [], ""
    for mot in mots:
        test = (cur + " " + mot).strip()
        if cur and (cur.endswith("-") or draw.textlength(test, font=f) > w):
            lignes.append(cur)
            cur = mot
        else:
            cur = test
    return (lignes + [cur] if cur else lignes) or [texte], f, int(taille * 1.1)


def image_flux(etapes, prelevements, largeur=1900) -> Image.Image:
    """Cartes d'étapes, frise du temps et badges ; noms des échantillons sous les cartes."""
    f_n, f_s, f_b, f_bs = police(False, 21), police(False, 18), police(True, 21), police(True, 17)
    par_ligne, gap = 5, 26
    cw = (largeur - (par_ligne - 1) * gap) // par_ligne
    par_apres: dict[int, list] = {}
    for p in prelevements:
        par_apres.setdefault(p["apres_etape"], []).append(p)
    tmp = ImageDraw.Draw(Image.new("RGB", (10, 10)))
    cartes = []
    for e in etapes:
        ls = []
        for r in (e.get("reactifs_detail") or [{"nom": n, "quantite": ""} for n in e.get("reactifs", [])]):
            if r.get("nom"):
                ls += [(t, f_b, "#14201D") for t in _lignes(tmp, r["nom"] + (" " + r["quantite"] if r.get("quantite") else ""), f_b, cw - 28)]
        cond = " · ".join(x for x in [conditions(e["parametres"])] + [f"{a['nom']} {a.get('valeur', '')}".strip() for a in e.get("autres", [])] if x)
        if cond:
            ls += [(t, f_n, "#0F6B5C") for t in _lignes(tmp, cond, f_n, cw - 28)]
        if e.get("equipement"):
            ls += [(t, f_s, "#566360") for t in _lignes(tmp, ", ".join(e["equipement"]), f_s, cw - 28)]
        tl, tfont, pas = _titre(tmp, e["operation"][:1].upper() + e["operation"][1:], cw - 72)
        extra = pas * (len(tl) - 1)
        ps = par_apres.get(e["n"], [])
        cartes.append({"e": e, "ls": ls, "ps": ps, "tl": tl, "tfont": tfont, "pas": pas, "extra": extra, "h": 70 + extra + 30 * len(ls) + 90})
    nb = max(1, -(-len(etapes) // par_ligne))
    hauteurs = []
    for r in range(nb):
        grp = cartes[r * par_ligne:(r + 1) * par_ligne] or [{"h": 120, "ps": []}]
        hauteurs.append(max(c["h"] for c in grp) + 64 * max(len(c["ps"]) for c in grp) + 50)
    haut_entree = 70
    img = Image.new("RGB", (largeur, haut_entree + sum(hauteurs) + 50), "white")
    d = ImageDraw.Draw(img)

    def pastille(x, y, p):
        l1 = f"{p['id']} · {fmt_vol(p['volume_ml'] or 0)} mL"
        l2 = moment_texte(p)
        w = int(max(d.textlength(l1, font=f_bs), d.textlength(l2, font=f_s)) + 64)
        d.rounded_rectangle([x, y, x + w, y + 54], 8, fill="#FDF1DC", outline="#8A4B00", width=2)
        d.ellipse([x + 8, y + 12, x + 38, y + 42], fill="#8A4B00")
        d.text((x + 23, y + 27), code_badge(p["id"]), font=f_bs, fill="white", anchor="mm")
        d.text((x + 48, y + 6), l1, font=f_bs, fill="#8A4B00")
        d.text((x + 48, y + 30), l2, font=f_s, fill="#8A4B00")
        return w

    d.text((0, 14), "ENTRÉE", font=f_bs, fill="#566360")
    px = d.textlength("ENTRÉE", font=f_bs) + 18
    for p in par_apres.get(0, []):
        px += pastille(px, 4, p) + 12
    y = haut_entree
    for r in range(nb):
        grp = cartes[r * par_ligne:(r + 1) * par_ligne]
        if not grp:
            break
        hc = max(c["h"] for c in grp)
        for k, c in enumerate(grp):
            x = k * (cw + gap)
            e = c["e"]
            d.rounded_rectangle([x, y, x + cw, y + hc], 12, fill="white", outline="#BFCBC7", width=3)
            d.ellipse([x + 12, y + 12, x + 48, y + 48], fill="#0F6B5C")
            d.text((x + 30, y + 30), str(e["n"]), font=f_bs, fill="white", anchor="mm")
            for k_l, lg in enumerate(c["tl"]):
                d.text((x + 60, y + 14 + k_l * c["pas"]), lg, font=c["tfont"], fill="#14201D")
            ly = y + 62 + c["extra"]
            for t, font, coul in c["ls"]:
                d.text((x + 14, ly), t, font=font, fill=coul)
                ly += 30
            dur = e["parametres"].get("duree_min")
            if dur or c["ps"]:
                fy = y + hc - 66
                x0, x1 = x + 22, x + cw - 22
                d.line([x0, fy + 12, x1, fy + 12], fill="#C9D3CF", width=5)
                d.text((x0, fy + 36), "0", font=f_s, fill="#566360")
                dl = fmt_dur(dur) if dur else "fin"
                d.text((x1 - d.textlength(dl, font=f_s), fy + 36), dl, font=f_s, fill="#566360")
                for p in c["ps"]:
                    pct = 1.0
                    if p.get("moment") == "pendant":
                        pct = min(1, max(0, p["temps_min"] / dur)) if dur and p.get("temps_min") is not None else 0.5
                    bx = x0 + (x1 - x0) * pct
                    d.ellipse([bx - 15, fy - 3, bx + 15, fy + 27], fill="#8A4B00")
                    d.text((bx, fy + 12), code_badge(p["id"]), font=f_bs, fill="white", anchor="mm")
            if k < len(grp) - 1:
                ay = y + 30
                d.line([x + cw + 4, ay, x + cw + gap - 8, ay], fill="#566360", width=3)
                d.polygon([(x + cw + gap - 4, ay), (x + cw + gap - 14, ay - 7), (x + cw + gap - 14, ay + 7)], fill="#566360")
            for j, p in enumerate(c["ps"]):
                pastille(x, y + hc + 12 + j * 64, p)
        y += hauteurs[r]
    so = "SORTIE · produit prétraité"
    d.text((largeur - d.textlength(so, font=f_bs), y - 30), so, font=f_bs, fill="#566360")
    return img


def ecrire(par, texte, taille=None, gras=False, couleur=None):
    """Ajoute du texte ; les repères « [… à saisir] » sont surlignés en jaune."""
    for seg in REPERE.split(texte):
        if not seg:
            continue
        run = par.add_run(seg)
        run.bold = gras
        if taille:
            run.font.size = Pt(taille)
        if couleur:
            run.font.color.rgb = RGBColor.from_string(couleur)
        if REPERE.fullmatch(seg):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.font.color.rgb = RGBColor.from_string("7A4A00")


def fond(cell, couleur):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur)
    tc_pr.append(shd)


def manque(v):
    return str(v) if v not in (None, "") else "[à saisir]"


def ecrire_docx(data: dict, chemin) -> None:
    ctx, etapes, pre = data.get("contexte", {}), data["etapes"], data.get("prelevements", [])
    inv = {r["nom"].lower(): r for r in data.get("inventaire_reactifs", [])}
    doc = Document()
    doc.core_properties.title = data.get("titre", "Manip")
    doc.core_properties.author = "Mat & Méthode"
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Segoe UI")
    for nom, taille, coul in (("Title", 20, "14201D"), ("Heading 1", 14, TEAL), ("Heading 2", 12, TEAL), ("Heading 3", 10.5, "14201D")):
        st = doc.styles[nom]
        st.font.name, st.font.size, st.font.bold = "Segoe UI", Pt(taille), True
        st.font.color.rgb = RGBColor.from_string(coul)

    doc.add_paragraph(data.get("titre", "Manip"), style="Title")
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c0 = t.rows[0].cells[0]
    fond(c0, "E1F0EC")
    p0 = c0.paragraphs[0]
    ecrire(p0, "Donneur d’ordre : ", gras=True)
    ecrire(p0, f"{manque(ctx.get('donneur'))}     ")
    if ctx.get("type"):
        ecrire(p0, "Type : ", gras=True)
        ecrire(p0, f"{ctx['type']}     ")
    ecrire(p0, "Réf. demande : ", gras=True)
    ecrire(p0, f"{manque(ctx.get('ref'))}     ")
    ecrire(p0, "Date : ", gras=True)
    ecrire(p0, str(ctx.get("date", "")))
    if ctx.get("tags"):
        p1 = c0.add_paragraph()
        ecrire(p1, "Tags : ", gras=True)
        ecrire(p1, "  ".join("#" + x for x in ctx["tags"]))
    lanc, jr = data.get("lancee"), data.get("journal") or []
    p2 = c0.add_paragraph()
    ecrire(p2, "Version : ", gras=True)
    ecrire(p2, f"{data.get('version', 1)} · enregistrée le {fmt_dt(data.get('enregistree_le'))}"
           + (f" · lancée le {fmt_dt(lanc['le'])} par {lanc['par']}" if lanc else ""))
    if lanc and jr:
        ecrire(c0.add_paragraph(), f"Modifiée après lancement ({len(jr)})", gras=True, couleur=AMBER)

    doc.add_heading("Flux de prétraitement", level=1)
    buf = io.BytesIO()
    image_flux(etapes, pre).save(buf, format="PNG", dpi=(250, 250))
    buf.seek(0)
    doc.add_picture(buf, width=Cm(17))
    doc.inline_shapes[-1]._inline.docPr.set("descr", "Schéma du flux : une carte par étape, badges des échantillons sur la frise du temps")
    vd = ctx.get("volumeDepart")
    if vd is not None:
        pris = sum(p["volume_ml"] or 0 for p in pre)
        reste = vd - pris
        ecrire(doc.add_paragraph(), f"Volume : {fmt_vol(vd)} mL au départ, {fmt_vol(pris)} mL prélevés, "
               + (f"{fmt_vol(reste)} mL restants. " if reste >= 0 else f"dépassement de {fmt_vol(-reste)} mL. ")
               + "Chaque badge marque un prélèvement : sa place sur la frise donne le moment (bord droit = sortie de l’étape).", 8.5, couleur=MUTED)

    doc.add_heading("Matériel et méthodes", level=1)
    for ss in data.get("sous_sections", []):
        if ss.get("texte"):
            doc.add_heading(ss["titre"], level=3)
            ecrire(doc.add_paragraph(), ss["texte"])

    doc.add_heading("Détail par phase du flux", level=1)
    for i, ph in enumerate(data.get("phases", [])):
        if i == 0:
            entete = f"{ctx.get('matiere') or ''} (lot {manque(ctx.get('lot'))})" + (f" · {fmt_vol(vd)} mL" if vd is not None else "")
        else:
            entete = (ph.get("conditions") or "conditions à compléter") + (" — " + ", ".join(ph["materiel"]) if ph.get("materiel") else "")
        lignes = [(ph["titre"], entete)]
        if i:
            lignes.append(("Réactifs", "\n".join(
                f"{r['nom']}{' · ' + r['quantite'] if r.get('quantite') else ''} — n° Scirius {manque(r.get('scirius'))}, "
                f"emplacement {manque(r.get('emplacement'))}, fournisseur {manque(r.get('fournisseur'))}, lot {manque(r.get('lot'))}"
                for r in ph["reactifs"]) or "aucun"))
        lignes.append(("Échantillons", "\n".join(
            f"{x['id']} · {fmt_vol(x['volume'] or 0)} mL · {x['moment']} · {x.get('labo', '')} · {', '.join(x.get('analyses', [])) or 'aucune analyse'} — "
            f"n° Scirius {manque(x.get('scirius'))}, n° Aramis {manque(x.get('aramis'))}" for x in ph["echantillons"]) or "aucun"))
        tab = doc.add_table(rows=len(lignes), cols=2)
        tab.style = "Table Grid"
        tab.autofit = False
        for k, (lab, txt) in enumerate(lignes):
            c1, c2 = tab.rows[k].cells
            c1.width, c2.width = Cm(3), Cm(14)
            ecrire(c1.paragraphs[0], lab, 9, gras=True, couleur=TEAL if k == 0 else MUTED)
            ecrire(c2.paragraphs[0], txt, 9)
            if k == 0:
                fond(c1, "E1F0EC")
                fond(c2, "E1F0EC")
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def tableau(titre, entetes, lignes, largeurs):
        doc.add_heading(titre, level=2)
        tb = doc.add_table(rows=1, cols=len(entetes))
        tb.style = "Table Grid"
        tb.autofit = False
        for k, h in enumerate(entetes):
            cell = tb.rows[0].cells[k]
            cell.width = Cm(largeurs[k])
            fond(cell, "E1F0EC")
            ecrire(cell.paragraphs[0], h, 8.5, gras=True)
        for ligne in lignes:
            cells = tb.add_row().cells
            for k, v in enumerate(ligne):
                cells[k].width = Cm(largeurs[k])
                ecrire(cells[k].paragraphs[0], manque(v), 8.5)

    noms: list[tuple[str, list[int]]] = []
    for e in etapes:
        for r in e.get("reactifs_detail") or [{"nom": n} for n in e.get("reactifs", [])]:
            if not r.get("nom"):
                continue
            x = next((y for y in noms if y[0].lower() == r["nom"].lower()), None)
            if x is None:
                noms.append((r["nom"], [e["n"]]))
            elif e["n"] not in x[1]:
                x[1].append(e["n"])
    if noms:
        tableau("Réactifs", ["Réactif", "N° Scirius", "Emplacement", "Fournisseur", "Lot", "Étape"],
                [[nm] + [inv.get(nm.lower(), {}).get(k) for k in ("scirius", "emplacement", "fournisseur", "lot")] + [", ".join(map(str, st))]
                 for nm, st in noms], [3.3, 2.2, 4.6, 2.8, 1.8, 2.3])
    if pre:
        tableau("Échantillons", ["Éch.", "Étape", "Moment", "Volume", "N° Scirius", "N° Aramis", "Laboratoire", "Analyses"],
                [[p["id"], "entrée" if not p["apres_etape"] else f"{p['apres_etape']} · {etapes[p['apres_etape'] - 1]['operation']}", moment_texte(p),
                  f"{fmt_vol(p['volume_ml'] or 0)} mL", p.get("scirius"), p.get("aramis"), p.get("labo"), ", ".join(p.get("analyses", []))] for p in pre],
                [1.2, 3.0, 2.4, 1.7, 1.9, 1.9, 2.0, 2.9])
    if lanc:
        if jr:
            tableau("Modifications après lancement", ["Date et heure", "Par", "Modification"],
                    [[fmt_dt(j["t"]), j["par"], j["txt"]] for j in jr], [3.6, 2.6, 10.8])
        else:
            doc.add_heading("Modifications après lancement", level=2)
            ecrire(doc.add_paragraph(), "Aucune modification depuis le lancement.")

    ecrire(sec.footer.paragraphs[0], "Document généré automatiquement - à relire avant diffusion", 8, couleur=MUTED)
    doc.save(str(chemin))
