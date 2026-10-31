"""Lecture du fichier « Suivi enregistrements scirius.xlsx » (SharePoint du magasin), comme le projet
Fractionnement : n° Scirius (Stock Item), nom du produit, lot (Batch), fournisseur, stockage…

Ordre des sources, seul un fichier vérifié (bonnes colonnes) est utilisé :
  1. fichier choisi à la main dans les Réglages ;
  2. fichier synchronisé dans le OneDrive / SharePoint de la personne connectée (dossier officiel d'abord) ;
  3. ouverture directe du lien SharePoint par Excel (Windows, pywin32), copie dans donnees\\cache ;
  4. dernière copie partagée donnees\\cache\\Cache_Scirius.xlsx.
Les lignes lues sont gardées dans donnees\\cache\\Cache_Scirius_lignes.json (30 min, ou tant que le fichier
synchronisé n'a pas changé) pour ne pas relire le classeur à chaque recherche.
"""
from __future__ import annotations

import json
import os
import re
import shutil
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Callable, Optional

FICHIER = "Suivi enregistrements scirius.xlsx"
URL_DIRECTE = ("https://totalworkplace.sharepoint.com/sites/MS-SMR-RECH-CRES-MAGASIN/"
               "Documents%20partages/Suivi%20enregistrement%20%C3%A9chantillons/"
               "Suivi%20enregistrements%20scirius.xlsx")
FEUILLE = "Suivi General"
LIGNE_ENTETE = 8
OBLIGATOIRES = ("date", "stock_item", "batch", "product", "dest")
CACHE_MINUTES = 30
ALIASES = {
    "date": ["date reception", "date de reception", "date recept"],
    "supplier": ["fournisseur"],
    "desired": ["product name souhaite si creation", "product name souhaite"],
    "stock_item": ["stock item", "stockitem"],
    "storage": ["stockage"],
    "nb": ["nb", "nombre"],
    "qty": ["quantite"],
    "batch": ["batch lot", "n batch lot", "batch", "lot"],
    "product": ["nom de produit", "nom du produit", "nom produit"],
    "dest": ["destinataire", "nom destinataire"],
    "equipe": ["equipe", "equipe bio", "team"],
}
COLONNES = [("date", "Date réception"), ("stock_item", "N° Scirius (Stock Item)"), ("product", "Nom du produit"),
            ("batch", "Lot"), ("supplier", "Fournisseur"), ("storage", "Stockage"), ("qty", "Quantité"), ("dest", "Destinataire")]


def texte(v) -> str:
    if v is None:
        return ""
    t = str(v)
    for ch in ("\xa0", "​", "‎", "‏", "﻿", "\r", "\n", "\t"):
        t = t.replace(ch, " ")
    return " ".join(t.split())


def cle_entete(v) -> str:
    t = unicodedata.normalize("NFKD", texte(v).lower()).encode("ascii", "ignore").decode("ascii")
    return " ".join(re.sub(r"[^a-z0-9]+", " ", t).split())


def racines_onedrive() -> list[Path]:
    """Dossiers OneDrive / SharePoint synchronisés de la personne connectée sur ce poste."""
    profil = Path(os.environ.get("USERPROFILE", str(Path.home())))
    out: list[Path] = []
    for v in (os.environ.get("OneDriveCommercial"), os.environ.get("OneDrive"),
              str(profil / "OneDrive - TotalEnergies"), str(profil / "TotalEnergies")):
        if v and Path(v).is_dir() and Path(v) not in out:
            out.append(Path(v))
    return out


def fichiers_synchronises(racines: list[Path] | None = None, profondeur: int = 6) -> list[Path]:
    """Fichiers Scirius trouvés dans les dossiers synchronisés : l'officiel d'abord, puis du plus récent au plus ancien."""
    trouves = []
    for racine in racines if racines is not None else racines_onedrive():
        base = len(racine.parts)
        for courant, dossiers, fichiers in os.walk(racine):
            if len(Path(courant).parts) - base >= profondeur:
                dossiers[:] = []
            dossiers[:] = [d for d in dossiers if not d.startswith((".", "~"))]
            trouves += [Path(courant) / f for f in fichiers if f.lower() == FICHIER.lower()]

    def officiel(p):
        k = cle_entete(str(p))
        return "suivi enregistrement echantillons" in k or "magasin" in k

    def date_modif(p):
        try:
            return p.stat().st_mtime
        except OSError:
            return 0
    return sorted(dict.fromkeys(trouves), key=lambda p: (not officiel(p), -date_modif(p)))


def ouvrir_avec_excel(url: str, cible: Path) -> tuple[Optional[Path], str]:
    """Ouvre le lien SharePoint avec Excel (compte Office de la personne connectée) et enregistre une copie."""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        return None, "Excel n’est pas accessible depuis l’application (module pywin32 absent)"
    excel = classeur = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        classeur = excel.Workbooks.Open(url, 0, True)
        classeur.SaveCopyAs(str(cible))
        return cible, ""
    except Exception as ex:
        return None, str(ex)
    finally:
        try:
            if classeur is not None:
                classeur.Close(SaveChanges=False)
            if excel is not None:
                excel.Quit()
        except Exception:
            pass


def associer_colonnes(valeurs) -> dict:
    """Associe chaque champ à une colonne : intitulé exact, sinon mots inclus (cellules courtes seulement)."""
    cellules = []
    for i, v in enumerate(valeurs or ()):
        t = cle_entete(v) if v is not None else ""
        if t:
            cellules.append((t, set(t.split()), i))
    trouve, pris = {}, set()
    for cle, noms in ALIASES.items():
        trouve[cle] = None
        for nom in noms:
            hit = next((i for t, _, i in cellules if t == nom and i not in pris), None)
            if hit is None:
                mots = set(nom.split())
                hit = next((i for t, tk, i in cellules if len(tk) <= 6 and mots.issubset(tk) and i not in pris), None)
            if hit is not None:
                trouve[cle] = hit
                pris.add(hit)
                break
    return trouve


def inspecter(chemin: Path) -> dict:
    """Vérifie que c'est bien le fichier Scirius (colonnes obligatoires) et repère la ligne d'en-tête."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(chemin, read_only=True, data_only=True)
    except Exception as ex:
        return {"ok": False, "raison": f"illisible ({ex})"}

    def ligne(ws, n):
        try:
            return next(iter(ws.iter_rows(min_row=n, max_row=n, max_col=100, values_only=True)))
        except StopIteration:
            return ()

    def score(cols):
        return sum(1 for k in OBLIGATOIRES if cols.get(k) is not None), sum(1 for v in cols.values() if v is not None)

    voulue = cle_entete(FEUILLE)
    feuilles = [ws for ws in wb.worksheets if cle_entete(ws.title) == voulue]
    feuilles += [ws for ws in wb.worksheets if ws not in feuilles]
    meilleur = None
    for k, ws in enumerate(feuilles):
        cols = associer_colonnes(ligne(ws, LIGNE_ENTETE))
        if score(cols)[0] == len(OBLIGATOIRES):
            meilleur = (score(cols), 0, ws, LIGNE_ENTETE, cols)
            break
        for n, valeurs in enumerate(ws.iter_rows(min_row=1, max_row=40, max_col=100, values_only=True), start=1):
            cols = associer_colonnes(valeurs)
            cand = (score(cols), -(abs(n - LIGNE_ENTETE) + 100 * k), ws, n, cols)
            if meilleur is None or cand[:2] > meilleur[:2]:
                meilleur = cand
        if meilleur and meilleur[0][0] == len(OBLIGATOIRES):
            break
    if meilleur is None or meilleur[0][0] < len(OBLIGATOIRES):
        wb.close()
        manque = [k for k in OBLIGATOIRES if not meilleur or meilleur[4].get(k) is None]
        return {"ok": False, "raison": f"colonnes manquantes ({', '.join(manque)}) : ancienne copie ou autre fichier"}
    return {"ok": True, "classeur": wb, "feuille": meilleur[2], "entete": meilleur[3], "colonnes": meilleur[4]}


def lire(info: dict) -> list[dict]:
    """Lignes du fichier (dictionnaires avec les clés d'ALIASES, plus « excel_row »)."""
    ws, entete, cols = info["feuille"], info["entete"], info["colonnes"]
    presentes = [i for i in cols.values() if i is not None]
    max_col = max(presentes) + 1
    lignes, vides = [], 0
    for n, valeurs in enumerate(ws.iter_rows(min_row=entete + 1, max_col=max_col, values_only=True), start=entete + 1):
        if not any(v is not None for v in valeurs):
            vides += 1
            if vides > 300:
                break
            continue
        vides = 0
        valeurs = tuple(valeurs) + (None,) * max(0, max_col - len(valeurs))
        item = {"excel_row": n}
        for cle, i in cols.items():
            v = valeurs[i] if i is not None else None
            if cle == "date" and isinstance(v, (datetime, date)):
                item[cle] = (v.date() if isinstance(v, datetime) else v).strftime("%d/%m/%Y")
            else:
                item[cle] = texte(v)
        lignes.append(item)
    info["classeur"].close()
    return lignes


def _date_tri(s: str):
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt)
        except (ValueError, TypeError):
            continue
    return datetime.min


def chercher(lignes: list[dict], requete: str, limite: int = 300) -> list[dict]:
    """Toutes les lignes contenant tous les mots de la requête (sans casse ni accents), les plus récentes d'abord."""
    mots = cle_entete(requete).split()
    out = []
    for ln in lignes:
        tout = cle_entete(" ".join(str(ln.get(k, "")) for k, _ in COLONNES))
        if all(m in tout for m in mots):
            out.append(ln)
    out.sort(key=lambda ln: _date_tri(ln.get("date", "")), reverse=True)
    return out[:limite]


@dataclass
class Lecture:
    lignes: list[dict] = field(default_factory=list)
    origine: str = ""
    ecartes: list[str] = field(default_factory=list)


class SourceScirius:
    """Trouve, vérifie et lit le fichier Scirius ; mémorise les lignes dans le dossier cache partagé."""

    def __init__(self, cache: Path, fichier_choisi: str = "", utilisateur: str = "", racines: list[Path] | None = None,
                 excel: Callable[[str, Path], tuple[Optional[Path], str]] = ouvrir_avec_excel):
        self.cache = Path(cache)
        self.fichier_choisi = fichier_choisi
        self.utilisateur = utilisateur
        self.racines = racines
        self.excel = excel

    @property
    def copie(self) -> Path:
        return self.cache / "Cache_Scirius.xlsx"

    @property
    def f_lignes(self) -> Path:
        return self.cache / "Cache_Scirius_lignes.json"

    @staticmethod
    def _signature(p) -> list | None:
        try:
            st = Path(p).stat()
            return [st.st_mtime, st.st_size]
        except (OSError, TypeError):
            return None

    def _cache_valable(self) -> Optional[dict]:
        try:
            d = json.loads(self.f_lignes.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return None
        if not isinstance(d, dict) or not d.get("lignes"):
            return None
        src = d.get("source") or ""
        sig = self._signature(src) if src and not str(src).startswith(str(self.cache)) else None
        if sig is not None:
            return d if sig == d.get("signature") else None
        age = datetime.now().timestamp() - float(d.get("lu_le") or 0)
        return d if 0 <= age < CACHE_MINUTES * 60 else None

    def _memoriser(self, lignes, source: Path, origine: str) -> None:
        try:
            self.cache.mkdir(parents=True, exist_ok=True)
            d = {"lignes": lignes, "lu_le": datetime.now().timestamp(), "lu_par": self.utilisateur, "origine": origine,
                 "source": str(source), "signature": self._signature(source)}
            tmp = self.f_lignes.with_name(self.f_lignes.name + ".tmp")
            tmp.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
            os.replace(tmp, self.f_lignes)
        except OSError:
            pass

    def _copier(self, source: Path) -> None:
        try:
            self.cache.mkdir(parents=True, exist_ok=True)
            tmp = self.cache / f"Cache_Scirius_{self.utilisateur or 'tmp'}.copie.tmp.xlsx"
            shutil.copy2(source, tmp)
            os.replace(tmp, self.copie)
        except OSError:
            pass

    def lire(self, forcer: bool = False, etape: Callable[[str], None] = lambda t: None) -> Lecture:
        if not forcer:
            d = self._cache_valable()
            if d:
                lu = datetime.fromtimestamp(float(d["lu_le"]))
                return Lecture(d["lignes"], f"{d.get('origine') or 'lecture précédente'}, lu le {lu:%d/%m à %H:%M}")
        res = Lecture()
        essayes = set()

        def essai(chemin, origine, copier=True) -> bool:
            chemin = Path(chemin)
            if str(chemin).lower() in essayes or not chemin.exists():
                return False
            essayes.add(str(chemin).lower())
            etape(f"Vérification : {chemin.name} ({origine})…")
            info = inspecter(chemin)
            if not info["ok"]:
                res.ecartes.append(f"{chemin} : {info['raison']}")
                return False
            etape("Lecture des lignes…")
            res.lignes = lire(info)
            res.origine = f"{origine}, lu le {datetime.now():%d/%m à %H:%M}"
            if copier:
                self._copier(chemin)
            self._memoriser(res.lignes, chemin, origine)
            return True

        if self.fichier_choisi and essai(self.fichier_choisi, "fichier choisi dans les Réglages"):
            return res
        etape("Recherche du fichier dans vos dossiers OneDrive / SharePoint…")
        for f in fichiers_synchronises(self.racines):
            if essai(f, "fichier synchronisé"):
                return res
        etape("Ouverture du fichier SharePoint avec Excel (peut prendre une minute)…")
        self.cache.mkdir(parents=True, exist_ok=True)
        chemin, erreur = self.excel(URL_DIRECTE, self.cache / f"Cache_Scirius_{self.utilisateur or 'tmp'}.tmp.xlsx")
        if chemin is not None:
            if essai(chemin, "SharePoint (Excel)"):
                return res
        else:
            res.ecartes.append(f"SharePoint via Excel : {erreur}")
        if self.copie.exists():
            jour = datetime.fromtimestamp(self.copie.stat().st_mtime)
            if essai(self.copie, f"copie partagée du {jour:%d/%m/%Y à %H:%M}", copier=False):
                return res
        return res
