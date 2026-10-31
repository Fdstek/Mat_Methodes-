"""Exporte une manip en Word (.docx) : schéma du flux (image), texte « Matériel et méthodes »,
détail par phase du flux, tableaux Réactifs et Échantillons.

Usage :
  python3 export_word.py manip-pour-script-pdf.json --out manip.docx

Entrée : le JSON d'échange (voir CAHIER_DES_CHARGES.md §7), le même que pour export_pdf_flux.py.
Dépendances : python-docx, Pillow.
"""
import argparse
import io
import json
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

TEAL, AMBER, MUTED = "0F6B5C", "8A4B00", "566360"
POLICES = [  # (normal, gras) : première paire trouvée
    ("C:/Windows/Fonts/segoeui.ttf", "C:/Windows/Fonts/segoeuib.ttf"),
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
]


def charger_police(gras, taille):
    for normal, bold in POLICES:
        chemin = bold if gras else normal
        if os.path.exists(chemin):
            return ImageFont.truetype(chemin, taille)
    return ImageFont.load_default()


def duree(minutes):
    h, m = divmod(int(round(minutes)), 60)
    if h and m:
        return f"{h} h {m:02d}"
    return f"{h} h" if h else f"{m} min"


def conditions(p):
    morceaux = []
    if "temperature_c" in p:
        morceaux.append(f"{p['temperature_c']:g} °C")
    if "duree_min" in p:
        morceaux.append(duree(p["duree_min"]))
    if "agitation_trmin" in p:
        morceaux.append(f"{p['agitation_trmin']} tr/min")
    if p.get("vide"):
        morceaux.append("sous vide" + (f" {p['vide_mbar']:g} mbar" if p.get("vide_mbar") else ""))
    if "filtre_um" in p:
        morceaux.append(f"{p['filtre_um']:g} µm")
    return " · ".join(morceaux)


def moment_texte(p):
    m = p.get("moment") or ("avant" if p["apres_etape"] == 0 else "effluent")
    if m == "avant":
        return "avant traitement"
    if m == "pendant":
        return f"pendant · à {duree(p['temps_min'])}" if p.get("temps_min") is not None else "pendant la réaction"
    return "effluent"


def code_badge(ident):
    m = re.match(r"^\D*(\d+)(.*)$", ident)
    return (m.group(2) or m.group(1))[:2] if m else ident[:2]


# ---------------------------------------------------------------- image du flux (Pillow)
def image_flux(etapes, prelevements, largeur=1900):
    """Dessine le flux : cartes d'étapes, frise du temps, badges ; noms des échantillons sous les cartes."""
    f_t, f_n, f_s = charger_police(True, 26), charger_police(False, 21), charger_police(False, 18)
    f_b, f_bs = charger_police(True, 21), charger_police(True, 17)
    par_ligne, gap = 5, 26
    cw = (largeur - (par_ligne - 1) * gap) // par_ligne
    par_apres = {}
    for p in prelevements:
        par_apres.setdefault(p["apres_etape"], []).append(p)

    def lignes(draw, texte, font, w):
        mots, out, cur = texte.split(), [], ""
        for mot in mots:
            test = (cur + " " + mot).strip()
            if draw.textlength(test, font=font) > w and cur:
                out.append(cur)
                cur = mot
            else:
                cur = test
        return out + ([cur] if cur else [])

    tmp = ImageDraw.Draw(Image.new("RGB", (10, 10)))
    cartes = []
    for e in etapes:
        ls = []
        for r in (e.get("reactifs_detail") or [{"nom": n, "quantite": ""} for n in e.get("reactifs", [])]):
            ls += [(t, f_b, "#14201D") for t in lignes(tmp, r["nom"] + (" " + r["quantite"] if r.get("quantite") else ""), f_b, cw - 28)]
        if conditions(e["parametres"]):
            ls += [(t, f_n, "#0F6B5C") for t in lignes(tmp, conditions(e["parametres"]), f_n, cw - 28)]
        if e.get("equipement"):
            ls += [(t, f_s, "#566360") for t in lignes(tmp, ", ".join(e["equipement"]), f_s, cw - 28)]
        ps = par_apres.get(e["n"], [])
        tfont, tl = f_t, lignes(tmp, e["operation"].capitalize(), f_t, cw - 72)
        if any(tmp.textlength(m, font=f_t) > cw - 72 for m in e["operation"].split()):
            tfont = charger_police(True, 20)
            tl = lignes(tmp, e["operation"].capitalize(), tfont, cw - 72)
        extra = 28 * (len(tl) - 1)
        cartes.append({"e": e, "ls": ls, "ps": ps, "tl": tl, "tfont": tfont, "extra": extra, "h": 70 + extra + 30 * len(ls) + 90})
    nb_lignes = max(1, -(-len(etapes) // par_ligne))
    hauteurs = []
    for r in range(nb_lignes):
        grp = cartes[r * par_ligne:(r + 1) * par_ligne]
        hauteurs.append(max(c["h"] for c in grp) + 64 * max(len(c["ps"]) for c in grp) + 50)
    haut_entree = 70
    img = Image.new("RGB", (largeur, haut_entree + sum(hauteurs) + 50), "white")
    d = ImageDraw.Draw(img)

    def pastille(x, y, p):
        l1 = f"{p['id']} · {p['volume_ml']:g} mL"
        l2 = moment_texte(p)
        w = int(max(d.textlength(l1, font=f_bs), d.textlength(l2, font=f_s)) + 64)
        d.rounded_rectangle([x, y, x + w, y + 54], 8, fill="#FDF1DC", outline="#8A4B00", width=2)
        d.ellipse([x + 8, y + 12, x + 38, y + 42], fill="#8A4B00")
        cb = code_badge(p["id"])
        d.text((x + 23, y + 27), cb, font=f_bs, fill="white", anchor="mm")
        d.text((x + 48, y + 6), l1, font=f_bs, fill="#8A4B00")
        d.text((x + 48, y + 30), l2, font=f_s, fill="#8A4B00")
        return w

    d.text((0, 14), "ENTRÉE", font=f_bs, fill="#566360")
    px = d.textlength("ENTRÉE", font=f_bs) + 18
    for p in par_apres.get(0, []):
        px += pastille(px, 4, p) + 12
    y = haut_entree
    for r in range(nb_lignes):
        grp = cartes[r * par_ligne:(r + 1) * par_ligne]
        hc = max(c["h"] for c in grp)
        for k, c in enumerate(grp):
            x = k * (cw + gap)
            e = c["e"]
            d.rounded_rectangle([x, y, x + cw, y + hc], 12, fill="white", outline="#BFCBC7", width=3)
            d.ellipse([x + 12, y + 12, x + 48, y + 48], fill="#0F6B5C")
            d.text((x + 30, y + 30), str(e["n"]), font=f_bs, fill="white", anchor="mm")
            for k_l, lg in enumerate(c["tl"]):
                d.text((x + 60, y + 14 + k_l * 28), lg, font=c["tfont"], fill="#14201D")
            ly = y + 62 + c["extra"]
            for t, font, coul in c["ls"]:
                d.text((x + 14, ly), t, font=font, fill=coul)
                ly += 30
            dur = e["parametres"].get("duree_min")
            if dur or c["ps"]:
                fy = y + hc - 66
                x0, x1 = x + 22, x + cw - 22
                d.line([x0, fy + 12, x1, fy + 12], fill="#C9D3CF", width=5)
                d.text((x0, fy + 36), "0", font=f_s, fill="#566360")
                dl = duree(dur) if dur else "fin"
                d.text((x1 - d.textlength(dl, font=f_s), fy + 36), dl, font=f_s, fill="#566360")
                for p in c["ps"]:
                    pct = 1.0
                    if p.get("moment") == "pendant":
                        pct = min(1, max(0, p["temps_min"] / dur)) if dur and p.get("temps_min") is not None else 0.5
                    bx = x0 + (x1 - x0) * pct
                    d.ellipse([bx - 15, fy - 3, bx + 15, fy + 27], fill="#8A4B00")
                    d.text((bx, fy + 12), code_badge(p["id"]), font=f_bs, fill="white", anchor="mm")
            if k < len(grp) - 1:
                ay = y + 30
                d.line([x + cw + 4, ay, x + cw + gap - 8, ay], fill="#566360", width=3)
                d.polygon([(x + cw + gap - 4, ay), (x + cw + gap - 14, ay - 7), (x + cw + gap - 14, ay + 7)], fill="#566360")
            for j, p in enumerate(c["ps"]):
                pastille(x, y + hc + 12 + j * 64, p)
        y += hauteurs[r]
    so = "SORTIE · produit prétraité"
    d.text((largeur - d.textlength(so, font=f_bs), y - 30), so, font=f_bs, fill="#566360")
    return img


# ---------------------------------------------------------------- Word
def surligner(par, texte, taille=None, gras=False, couleur=None):
    """Ajoute du texte ; les repères [xxx à saisir] sont surlignés en jaune."""
    for seg in re.split(r"(\[[^\]]* à saisir\])", texte):
        if not seg:
            continue
        run = par.add_run(seg)
        run.bold = gras
        if taille:
            run.font.size = Pt(taille)
        if couleur:
            run.font.color.rgb = RGBColor.from_string(couleur)
        if re.fullmatch(r"\[[^\]]* à saisir\]", seg):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.font.color.rgb = RGBColor.from_string("7A4A00")


def remplir_cellule(cell, fond=None):
    if fond:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fond)
        tc_pr.append(shd)


def fmt_dt(iso):
    from datetime import datetime
    try:
        return datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone().strftime("%d/%m/%Y à %H:%M")
    except Exception:
        return str(iso or "")


def manque(v):
    return v if v not in (None, "") else "[à saisir]"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("json")
    ap.add_argument("--out", default="manip.docx")
    a = ap.parse_args()
    data = json.load(open(a.json, encoding="utf-8"))
    ctx, etapes, pre = data.get("contexte", {}), data["etapes"], data.get("prelevements", [])
    inv = {r["nom"].lower(): r for r in data.get("inventaire_reactifs", [])}

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Segoe UI")
    for nom, taille, coul in (("Title", 20, "14201D"), ("Heading 1", 14, TEAL), ("Heading 2", 12, TEAL), ("Heading 3", 10.5, "14201D")):
        st = doc.styles[nom]
        st.font.name, st.font.size, st.font.bold = "Segoe UI", Pt(taille), True
        st.font.color.rgb = RGBColor.from_string(coul)

    doc.add_paragraph(data.get("titre", "Manip"), style="Title")
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c0 = t.rows[0].cells[0]
    remplir_cellule(c0, "E1F0EC")
    p0 = c0.paragraphs[0]
    surligner(p0, "Donneur d’ordre : ", gras=True)
    surligner(p0, f"{manque(ctx.get('donneur'))}     ")
    if ctx.get("type"):
        surligner(p0, "Type : ", gras=True)
        surligner(p0, f"{ctx['type']}     ")
    surligner(p0, "Réf. demande : ", gras=True)
    surligner(p0, f"{manque(ctx.get('ref'))}     ")
    surligner(p0, "Date : ", gras=True)
    surligner(p0, str(ctx.get("date", "")))
    if ctx.get("tags"):
        p1 = c0.add_paragraph()
        surligner(p1, "Tags : ", gras=True)
        surligner(p1, "  ".join("#" + t for t in ctx["tags"]))
    lanc, jr = data.get("lancee"), data.get("journal") or []
    if data.get("version"):
        p2 = c0.add_paragraph()
        surligner(p2, "Version : ", gras=True)
        surligner(p2, f"{data['version']} · enregistrée le {fmt_dt(data.get('enregistree_le'))}"
                  + (f" · lancée le {fmt_dt(lanc['le'])} par {lanc['par']}" if lanc else ""))
        if lanc and jr:
            p3 = c0.add_paragraph()
            surligner(p3, f"Modifiée après lancement ({len(jr)})", gras=True, couleur=AMBER)

    doc.add_heading("Flux de prétraitement", level=1)
    buf = io.BytesIO()
    image_flux(etapes, pre).save(buf, format="PNG", dpi=(250, 250))
    buf.seek(0)
    doc.add_picture(buf, width=Cm(17))
    doc.paragraphs[-1].alignment = WD_TABLE_ALIGNMENT.LEFT
    pic = doc.inline_shapes[-1]
    pic._inline.docPr.set("descr", "Schéma du flux : une carte par étape, badges des échantillons sur la frise du temps")
    vd = ctx.get("volumeDepart")
    if vd:
        pris = sum(p["volume_ml"] for p in pre)
        surligner(doc.add_paragraph(), f"Volume : {vd:g} mL au départ, {pris:g} mL prélevés, {vd - pris:g} mL restants. "
                  "Chaque badge marque un prélèvement : sa place sur la frise donne le moment (bord droit = sortie de l’étape).", 8.5, couleur=MUTED)

    doc.add_heading("Matériel et méthodes", level=1)
    for ss in data.get("sous_sections", []):
        if ss.get("texte"):
            doc.add_heading(ss["titre"], level=3)
            surligner(doc.add_paragraph(), ss["texte"])

    doc.add_heading("Détail par phase du flux", level=1)
    phases = [(0, "Départ", "", [], [])] + [
        (e["n"], f"{e['n']} · {e['operation'].capitalize()}", conditions(e["parametres"]), e.get("equipement", []),
         e.get("reactifs_detail") or [{"nom": n, "quantite": ""} for n in e.get("reactifs", [])]) for e in etapes]
    for n, titre, cond, equip, reacs in phases:
        lignes = [("Départ" if not n else titre, f"{ctx.get('matiere', '')} (lot {manque(ctx.get('lot'))}) · {vd:g} mL" if not n and vd
                   else (cond or "conditions à compléter") + (" — " + ", ".join(equip) if equip else ""))]
        if n:
            lignes.append(("Réactifs", "\n".join(
                f"{r['nom']}{' · ' + r['quantite'] if r.get('quantite') else ''} — n° Scirius {manque(inv.get(r['nom'].lower(), {}).get('scirius'))}, "
                f"emplacement {manque(inv.get(r['nom'].lower(), {}).get('emplacement'))}, fournisseur {manque(inv.get(r['nom'].lower(), {}).get('fournisseur'))}, "
                f"lot {manque(inv.get(r['nom'].lower(), {}).get('lot'))}" for r in reacs) or "aucun"))
        ps = [p for p in pre if p["apres_etape"] == n]
        lignes.append(("Échantillons", "\n".join(
            f"{p['id']} · {p['volume_ml']:g} mL · {moment_texte(p)} · {p.get('labo', '')} · {', '.join(p.get('analyses', [])) or 'aucune analyse'} — "
            f"n° Scirius {manque(p.get('scirius'))}, n° Aramis {manque(p.get('aramis'))}" for p in ps) or "aucun"))
        tab = doc.add_table(rows=len(lignes), cols=2)
        tab.style = "Table Grid"
        tab.autofit = False
        tab.columns[0].width, tab.columns[1].width = Cm(3), Cm(14)
        for i, (lab, txt) in enumerate(lignes):
            c1, c2 = tab.rows[i].cells
            c1.width, c2.width = Cm(3), Cm(14)
            surligner(c1.paragraphs[0], lab, 9, gras=True, couleur=TEAL if i == 0 else MUTED)
            surligner(c2.paragraphs[0], txt, 9)
            if i == 0:
                remplir_cellule(c1, "E1F0EC")
                remplir_cellule(c2, "E1F0EC")
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def tableau(titre, entetes, lignes, largeurs):
        doc.add_heading(titre, level=2)
        tb = doc.add_table(rows=1, cols=len(entetes))
        tb.style = "Table Grid"
        tb.autofit = False
        for k, w in enumerate(largeurs):
            tb.columns[k].width = Cm(w)
        for k, h in enumerate(entetes):
            cell = tb.rows[0].cells[k]
            cell.width = Cm(largeurs[k])
            remplir_cellule(cell, "E1F0EC")
            surligner(cell.paragraphs[0], h, 8.5, gras=True)
        for ligne in lignes:
            cells = tb.add_row().cells
            for k, v in enumerate(ligne):
                cells[k].width = Cm(largeurs[k])
                surligner(cells[k].paragraphs[0], str(manque(v)), 8.5)

    noms = []
    for e in etapes:
        for r in (e.get("reactifs_detail") or [{"nom": n} for n in e.get("reactifs", [])]):
            if r["nom"].lower() not in [x.lower() for x, _ in noms]:
                noms.append((r["nom"], [e["n"]]))
            else:
                [s for x, s in noms if x.lower() == r["nom"].lower()][0].append(e["n"])
    if noms:
        tableau("Réactifs", ["Réactif", "N° Scirius", "Emplacement", "Fournisseur", "Lot", "Étape"],
                [[nm, inv.get(nm.lower(), {}).get("scirius"), inv.get(nm.lower(), {}).get("emplacement"), inv.get(nm.lower(), {}).get("fournisseur"),
                  inv.get(nm.lower(), {}).get("lot"), ", ".join(str(x) for x in st)] for nm, st in noms], [3.3, 2.2, 4.6, 2.8, 1.8, 2.3])
    if pre:
        tableau("Échantillons", ["Éch.", "Étape", "Moment", "Volume", "N° Scirius", "N° Aramis", "Laboratoire", "Analyses"],
                [[p["id"], "entrée" if not p["apres_etape"] else f"{p['apres_etape']} · {etapes[p['apres_etape'] - 1]['operation']}", moment_texte(p),
                  f"{p['volume_ml']:g} mL", p.get("scirius"), p.get("aramis"), p.get("labo"), ", ".join(p.get("analyses", []))] for p in pre],
                [1.2, 3.0, 2.4, 1.7, 1.9, 1.9, 2.0, 2.9])

    if lanc:
        if jr:
            tableau("Modifications après lancement", ["Date et heure", "Par", "Modification"],
                    [[fmt_dt(j["t"]), j["par"], j["txt"]] for j in jr], [3.6, 2.6, 10.8])
        else:
            doc.add_heading("Modifications après lancement", level=2)
            surligner(doc.add_paragraph(), "Aucune modification depuis le lancement.")

    pied = sec.footer.paragraphs[0]
    surligner(pied, "Document généré automatiquement - à relire avant diffusion", 8, couleur=MUTED)
    doc.save(a.out)
    print("Word écrit :", a.out)


if __name__ == "__main__":
    main()
